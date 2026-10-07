from docx import Document
from docx.oxml.ns import qn

from utils.markdown_docx import markdown_to_docx


def test_word_fallback_keeps_own_analysis_risk_table_and_source_url(tmp_path):
    source = tmp_path / "综合投研报告.md"
    source.write_text(
        "# 公司研究报告\n\n## 独立财务分析\n2025 年收入增长。\n\n"
        "## 风险评估\n\n| 风险 | 来源 |\n| --- | --- |\n"
        "| 原料价格 | [S1](https://example.com/report) |\n",
        encoding="utf-8",
    )
    output = markdown_to_docx(source, tmp_path / "综合投研报告.docx")
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "独立财务分析" in text
    assert "2025 年收入增长" in text
    assert "风险评估" in text
    assert "原料价格" in text
    assert "https://example.com/report" in text


def test_word_fallback_keeps_short_annual_metrics_as_tables(tmp_path):
    source = tmp_path / "自研报告.md"
    source.write_text(
        "# 财务分析\n\n| 年度 | 营收同比 | 毛利率 |\n"
        "| --- | ---: | ---: |\n| 2025 | 20.0% | 25.0% |\n",
        encoding="utf-8",
    )
    document = Document(markdown_to_docx(source, tmp_path / "自研报告.docx"))
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 1).text == "20.0%"
    assert document.tables[0].rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
