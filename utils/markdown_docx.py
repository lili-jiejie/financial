"""Portable Word fallback when Pandoc is unavailable.

This converter retains report text, source URLs, headings and tables. Pandoc
remains the preferred path for full Markdown styling and image handling.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _plain(text: str) -> str:
    text = re.sub(r"!?\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"^>\s*", "", text)
    return text.replace(r"\|", "|").replace("**", "").replace("`", "")


def _cells(line: str) -> list[str]:
    values = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [_plain(cell.strip()) for cell in values]


def markdown_to_docx(markdown_path: str | Path, output_path: str | Path) -> Path:
    source = Path(markdown_path)
    output = Path(output_path)
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Title"].font.size = Pt(18)
    title_properties = document.styles["Title"].element.pPr
    if title_properties is not None:
        for border in title_properties.findall(qn("w:pBdr")):
            title_properties.remove(border)
    source_links: list[tuple[str, str]] = []
    seen_urls: set[str] = set()
    for line in lines:
        for label, url in re.findall(r"(?<!!)\[([^\]]+)\]\((https?://[^)]+)\)", line):
            if url not in seen_urls:
                source_links.append((label, url))
                seen_urls.add(url)
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line == "---":
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\|?\s*:?-{3,}", lines[index + 1].strip()
        ):
            headers = _cells(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_cells(lines[index]))
                index += 1
            if "年度" in headers and "营业收入" in headers:
                table = document.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                table.autofit = True
                for cell, value in zip(table.rows[0].cells, headers):
                    cell.text = value
                for values in rows:
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = value
            else:
                # Long research quotes and URLs become unreadable in narrow
                # Word table columns. Keep each record together at page width.
                for values in rows:
                    paragraph = document.add_paragraph(style="List Bullet")
                    paragraph.paragraph_format.keep_together = True
                    for field, value in zip(headers, values):
                        if value:
                            paragraph.add_run(f"{field}：").bold = True
                            paragraph.add_run(value + "\n")
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            if level == 1:
                document.add_paragraph(_plain(heading.group(2)), style="Title")
            else:
                document.add_heading(_plain(heading.group(2)), level=min(level - 1, 4))
            index += 1
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image:
            image_path = (source.parent / image.group(2)).resolve()
            if (
                image_path.is_relative_to(source.parent.resolve())
                and image_path.is_file()
                and image_path.suffix.lower() in (".png", ".jpg", ".jpeg")
            ):
                document.add_picture(str(image_path))
            else:
                document.add_paragraph(_plain(line))
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            document.add_paragraph(_plain(bullet.group(1)), style="List Bullet")
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            document.add_paragraph(_plain(numbered.group(1)), style="List Number")
            index += 1
            continue
        document.add_paragraph(_plain(line))
        index += 1
    if source_links:
        document.add_heading("来源网址", level=1)
        for label, url in source_links:
            document.add_paragraph(f"{_plain(label)}：{url}")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output
